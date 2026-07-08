"""
main_window.py — Giao diện chính KathTTS Studio (PyQt6, Dark Theme)
"""

from __future__ import annotations

import sys
import tempfile
import winsound
from pathlib import Path
from typing import Optional

from PyQt6.QtCore import (
    Qt, QThread, QTimer, pyqtSignal, QPropertyAnimation,
    QEasingCurve, QSize, QSettings,
)
from PyQt6.QtGui import (
    QColor, QDragEnterEvent, QDropEvent, QFont,
    QLinearGradient, QPainter, QPalette,
)
from PyQt6.QtWidgets import (
    QApplication, QCheckBox, QComboBox, QFileDialog,
    QFrame, QHBoxLayout, QLabel, QLineEdit, QMainWindow,
    QMessageBox, QPlainTextEdit, QProgressBar, QPushButton,
    QScrollArea, QSizePolicy, QSplitter, QStackedWidget, QVBoxLayout, QWidget,
)

from app.ui.slide_sync_tab import SlideSyncTab

# ═══════════════════════════════════════════════════════════════════════════
#  STYLESHEET — Dark purple theme
# ═══════════════════════════════════════════════════════════════════════════

STYLESHEET = """
/* ── Base ── */
QMainWindow, QDialog { background-color: #0d1117; }

QWidget {
    background-color: #0d1117;
    color: #e6edf3;
    font-family: 'Segoe UI', 'Arial', sans-serif;
    font-size: 13px;
}

/* ── Cards ── */
QFrame#card {
    background-color: #161b22;
    border: 1px solid #30363d;
    border-radius: 10px;
    padding: 4px;
}

/* ── Buttons ── */
QPushButton {
    background-color: #21262d;
    color: #c9d1d9;
    border: 1px solid #30363d;
    border-radius: 6px;
    padding: 7px 16px;
    font-weight: 600;
    min-height: 28px;
}
QPushButton:hover  { background-color: #30363d; border-color: #6e7681; }
QPushButton:pressed{ background-color: #161b22; }
QPushButton:disabled { background-color: #161b22; color: #3d444d; border-color: #21262d; }

QPushButton#primary {
    background-color: #6d28d9;
    color: #ffffff;
    border: none;
    border-radius: 6px;
}
QPushButton#primary:hover   { background-color: #7c3aed; }
QPushButton#primary:pressed { background-color: #5b21b6; }
QPushButton#primary:disabled{ background-color: #1e2028; color: #3d444d; border: 1px solid #21262d; }

QPushButton#success {
    background-color: #1a7f37;
    color: #ffffff;
    border: none;
    border-radius: 8px;
    font-size: 14px;
    font-weight: 700;
    padding: 12px 20px;
    min-height: 44px;
}
QPushButton#success:hover   { background-color: #238636; }
QPushButton#success:pressed { background-color: #116329; }
QPushButton#success:disabled{ background-color: #1e2028; color: #3d444d; border: 1px solid #21262d; }

QPushButton#danger {
    background-color: #b91c1c;
    color: #ffffff;
    border: none;
    border-radius: 6px;
}
QPushButton#danger:hover { background-color: #dc2626; }

/* ── Inputs ── */
QPlainTextEdit {
    background-color: #0d1117;
    border: 1px solid #30363d;
    border-radius: 8px;
    color: #e6edf3;
    font-size: 14px;
    padding: 10px;
    line-height: 1.6;
    selection-background-color: #6d28d9;
}
QPlainTextEdit:focus { border-color: #6d28d9; }

QLineEdit {
    background-color: #161b22;
    border: 1px solid #30363d;
    border-radius: 6px;
    color: #e6edf3;
    padding: 6px 10px;
    min-height: 28px;
}
QLineEdit:focus { border-color: #6d28d9; }

QComboBox {
    background-color: #161b22;
    border: 1px solid #30363d;
    border-radius: 6px;
    color: #e6edf3;
    padding: 6px 10px;
    min-height: 28px;
}
QComboBox:hover  { border-color: #6d28d9; }
QComboBox:focus  { border-color: #6d28d9; }
QComboBox::drop-down { border: none; padding-right: 6px; }
QComboBox QAbstractItemView {
    background-color: #161b22;
    border: 1px solid #30363d;
    color: #e6edf3;
    selection-background-color: #6d28d9;
    outline: none;
}

/* ── Checkboxes ── */
QCheckBox { color: #c9d1d9; spacing: 8px; }
QCheckBox::indicator {
    width: 16px; height: 16px;
    border: 1px solid #484f58;
    border-radius: 4px;
    background-color: #21262d;
}
QCheckBox::indicator:checked {
    background-color: #6d28d9;
    border-color: #6d28d9;
    image: none;
}

/* ── Progress bars ── */
QProgressBar {
    background-color: #21262d;
    border: none;
    border-radius: 4px;
    height: 8px;
    text-align: center;
    color: transparent;
}
QProgressBar::chunk {
    background: qlineargradient(
        x1:0, y1:0, x2:1, y2:0,
        stop:0 #6d28d9, stop:1 #9333ea
    );
    border-radius: 4px;
}

/* ── Scrollbars ── */
QScrollBar:vertical {
    background-color: #0d1117;
    width: 8px;
    border-radius: 4px;
    margin: 0;
}
QScrollBar::handle:vertical {
    background-color: #30363d;
    border-radius: 4px;
    min-height: 24px;
}
QScrollBar::handle:vertical:hover { background-color: #484f58; }
QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical { height: 0; }

QScrollArea { border: none; }

/* ── Splitter ── */
QSplitter::handle { background-color: #21262d; }

/* ── Labels ── */
QLabel { background-color: transparent; }
QLabel#heading {
    font-size: 11px;
    font-weight: 700;
    color: #7d8590;
    letter-spacing: 0.08em;
}
QLabel#badge {
    background-color: #21262d;
    border: 1px solid #30363d;
    border-radius: 10px;
    color: #8b949e;
    font-size: 11px;
    padding: 2px 8px;
}
QLabel#badge-green {
    background-color: #0f4f1c;
    border: 1px solid #238636;
    border-radius: 10px;
    color: #3fb950;
    font-size: 11px;
    padding: 2px 8px;
}
QLabel#badge-orange {
    background-color: #3d1f00;
    border: 1px solid #f0883e;
    border-radius: 10px;
    color: #f0883e;
    font-size: 11px;
    padding: 2px 8px;
}

/* ── Step indicator ── */
QLabel#step-active {
    color: #c4b5fd;
    font-size: 12px;
    font-weight: 700;
    background: transparent;
}
QLabel#step-inactive {
    color: #484f58;
    font-size: 12px;
    background: transparent;
}
QLabel#step-arrow {
    color: #30363d;
    font-size: 12px;
    background: transparent;
    padding: 0 6px;
}
"""


# ═══════════════════════════════════════════════════════════════════════════
#  HELPERS
# ═══════════════════════════════════════════════════════════════════════════

def _read_script_file(path: str) -> str:
    """Đọc nội dung file văn bản (.txt hoặc .docx) và trả về string."""
    p = Path(path)
    if p.suffix.lower() == ".docx":
        try:
            import docx
            doc = docx.Document(str(p))
            return "\n".join(para.text for para in doc.paragraphs if para.text.strip())
        except ImportError:
            raise RuntimeError(
                "Cần thư viện python-docx để đọc file .docx\n"
                "Chạy: pip install python-docx"
            )
    else:
        # Thử UTF-8, fallback sang cp1252
        for enc in ("utf-8", "utf-8-sig", "cp1252"):
            try:
                return p.read_text(encoding=enc)
            except (UnicodeDecodeError, LookupError):
                continue
        return p.read_text(errors="replace")


# ═══════════════════════════════════════════════════════════════════════════
#  WORKER THREADS
# ═══════════════════════════════════════════════════════════════════════════

class DownloadThread(QThread):
    progress = pyqtSignal(int, str)
    finished = pyqtSignal(bool, str)   # success, error_message

    def __init__(self, engine, model_name: str):
        super().__init__()
        self.engine     = engine
        self.model_name = model_name

    def run(self):
        try:
            self.engine.download_model(self.model_name, self.progress.emit)
            self.finished.emit(True, "")
        except Exception as exc:
            self.finished.emit(False, str(exc))


class ExportThread(QThread):
    progress = pyqtSignal(int, str)
    finished = pyqtSignal(bool, str)   # success, error_message

    def __init__(self, pipeline, text, engine, speaker_id, output_path, use_whisper, speed=1.0):
        super().__init__()
        self.pipeline    = pipeline
        self.text        = text
        self.engine      = engine
        self.speaker_id  = speaker_id
        self.output_path = output_path
        self.use_whisper = use_whisper
        self.speed       = speed

    def run(self):
        try:
            self.pipeline.run(
                self.text,
                self.engine,
                self.speaker_id,
                self.output_path,
                self.progress.emit,
                self.use_whisper,
                speed=self.speed,
            )
            self.finished.emit(True, "")
        except Exception as exc:
            self.finished.emit(False, str(exc))


class PreviewThread(QThread):
    finished = pyqtSignal(bool, str, str)  # success, wav_path, error

    def __init__(self, engine, model_name, speaker_id, text, speed=1.0):
        super().__init__()
        self.engine      = engine
        self.model_name  = model_name
        self.speaker_id  = speaker_id
        self.text        = text
        self.speed       = speed
        self._tmp_wav    = ""

    def run(self):
        try:
            self.engine.load_model(self.model_name)
            tmp = tempfile.NamedTemporaryFile(suffix=".wav", delete=False)
            tmp.close()
            self.engine.synthesize(self.text, tmp.name, self.speaker_id, speed=self.speed)
            self.finished.emit(True, tmp.name, "")
        except Exception as exc:
            self.finished.emit(False, "", str(exc))


# ═══════════════════════════════════════════════════════════════════════════
#  DROP ZONE WIDGET
# ═══════════════════════════════════════════════════════════════════════════

class DropZoneWidget(QFrame):
    """Drag-and-drop zone để mở file .txt / .docx."""
    file_dropped = pyqtSignal(str)

    _STYLE_IDLE = """
        QFrame#dropZone {
            background-color: #161b22;
            border: 2px dashed #30363d;
            border-radius: 12px;
        }
    """
    _STYLE_HOVER = """
        QFrame#dropZone {
            background-color: #1a0e36;
            border: 2px dashed #6d28d9;
            border-radius: 12px;
        }
    """

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setObjectName("dropZone")
        self.setAcceptDrops(True)
        self.setMinimumHeight(130)
        self.setMaximumHeight(155)
        self.setStyleSheet(self._STYLE_IDLE)

        layout = QVBoxLayout(self)
        layout.setSpacing(6)
        layout.setAlignment(Qt.AlignmentFlag.AlignCenter)

        self._icon = QLabel("📂")
        self._icon.setStyleSheet("font-size: 36px; background:transparent;")
        self._icon.setAlignment(Qt.AlignmentFlag.AlignCenter)

        self._lbl = QLabel("Kéo & thả file  .docx  hoặc  .txt  vào đây")
        self._lbl.setStyleSheet("color: #7d8590; font-size: 13px; background:transparent;")
        self._lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)

        btn = QPushButton("Chọn file…")
        btn.setFixedWidth(110)
        btn.clicked.connect(self._open_dialog)

        layout.addWidget(self._icon)
        layout.addWidget(self._lbl)
        layout.addSpacing(4)
        layout.addWidget(btn, alignment=Qt.AlignmentFlag.AlignCenter)

    def _open_dialog(self):
        settings = QSettings("KathTTS", "KathSlideToVideoMaker")
        last_dir = settings.value("last_export_dir", None)
        if not last_dir:
            last_dir = str(Path.home() / "Downloads")
            if not Path(last_dir).exists():
                last_dir = str(Path.home())

        path, _ = QFileDialog.getOpenFileName(
            self, "Mở file văn bản", last_dir,
            "Văn bản (*.txt *.docx);;Tất cả (*.*)"
        )
        if path:
            self.file_dropped.emit(path)
            settings.setValue("last_export_dir", str(Path(path).parent))

    # ── Drag events ──────────────────────────────────────────────────────

    def dragEnterEvent(self, event: QDragEnterEvent):
        if event.mimeData().hasUrls():
            url = event.mimeData().urls()[0].toLocalFile()
            if url.lower().endswith((".txt", ".docx")):
                event.acceptProposedAction()
                self.setStyleSheet(self._STYLE_HOVER)
                self._lbl.setText("Thả file vào đây ✓")
                return
        event.ignore()

    def dragLeaveEvent(self, event):
        self.setStyleSheet(self._STYLE_IDLE)
        self._lbl.setText("Kéo & thả file  .docx  hoặc  .txt  vào đây")

    def dropEvent(self, event: QDropEvent):
        self.setStyleSheet(self._STYLE_IDLE)
        self._lbl.setText("Kéo & thả file  .docx  hoặc  .txt  vào đây")
        url = event.mimeData().urls()[0].toLocalFile()
        self.file_dropped.emit(url)


# ═══════════════════════════════════════════════════════════════════════════
#  MAIN WINDOW
# ═══════════════════════════════════════════════════════════════════════════

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("KathFlow Studio — Slide To Video")
        self.setMinimumSize(960, 640)
        self.resize(1180, 740)
        self.setStyleSheet(STYLESHEET)

        # ── Core objects ─────────────────────────────────────────────────
        from app.core.tts_engine import TTSEngine, VI_MALE_MODELS
        from app.core.mp3_exporter import ExportPipeline
        from app.core.file_reader import read_file

        self._engine        = TTSEngine()
        self._pipeline      = ExportPipeline()
        self._read_file     = read_file
        self._vi_models     = VI_MALE_MODELS

        self._download_thread: Optional[QThread] = None
        self._export_thread:   Optional[QThread] = None
        self._preview_thread:  Optional[QThread] = None

        # ── Build UI ─────────────────────────────────────────────────────
        self._build_ui()
        self._refresh_model_status()

    # ═══════════════════════════════════════════════════════════════════
    #  UI BUILD
    # ═══════════════════════════════════════════════════════════════════

    def _build_ui(self):
        root = QWidget()
        self.setCentralWidget(root)
        root_layout = QVBoxLayout(root)
        root_layout.setContentsMargins(0, 0, 0, 0)
        root_layout.setSpacing(0)

        # Shared header with step indicator
        root_layout.addWidget(self._build_header())

        # ── Stacked pages ─────────────────────────────────────────────
        self._stack = QStackedWidget()

        # Page 0: MP3 creation (existing layout)
        mp3_page = QWidget()
        mp3_layout = QVBoxLayout(mp3_page)
        mp3_layout.setContentsMargins(0, 0, 0, 0)
        mp3_layout.setSpacing(0)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.setHandleWidth(1)
        splitter.addWidget(self._build_left())
        splitter.addWidget(self._build_right())
        splitter.setSizes([580, 440])  # Tăng chiều rộng panel phải để tránh scrollbar ngang
        mp3_layout.addWidget(splitter, 1)
        mp3_layout.addWidget(self._build_footer())

        self._stack.addWidget(mp3_page)

        # Page 1: Slide sync tab
        self._slide_sync_tab = SlideSyncTab()
        self._slide_sync_tab.request_back.connect(self._go_to_mp3_tab)
        self._slide_sync_tab.request_export.connect(self._on_export_video_requested)
        self._stack.addWidget(self._slide_sync_tab)

        root_layout.addWidget(self._stack, 1)

    # ── Header ──────────────────────────────────────────────────────────

    def _build_header(self) -> QFrame:
        hdr = QFrame()
        hdr.setFixedHeight(56)
        hdr.setStyleSheet("""
            QFrame {
                background-color: #161b22;
                border-bottom: 1px solid #21262d;
            }
        """)
        lay = QHBoxLayout(hdr)
        lay.setContentsMargins(20, 0, 20, 0)

        title = QLabel("🎬  KathFlow Studio")
        title.setStyleSheet(
            "font-size: 17px; font-weight: 700; color: #e6edf3; background:transparent;"
        )
        lay.addWidget(title)
        lay.addSpacing(20)

        # ── Step indicator breadcrumb ──────────────────────────────────
        self._step1_lbl = QLabel("● Bước 1: Tạo MP3")
        self._step1_lbl.setObjectName("step-active")

        arrow1 = QLabel("→")
        arrow1.setObjectName("step-arrow")

        self._step2_lbl = QLabel("○ Bước 2: Đồng bộ Slide")
        self._step2_lbl.setObjectName("step-inactive")

        arrow2 = QLabel("→")
        arrow2.setObjectName("step-arrow")

        self._step3_lbl = QLabel("○ Bước 3: Xuất Video")
        self._step3_lbl.setObjectName("step-inactive")

        for w in [self._step1_lbl, arrow1, self._step2_lbl, arrow2, self._step3_lbl]:
            lay.addWidget(w)

        lay.addStretch()
        return hdr

    # ── Left panel ───────────────────────────────────────────────────────

    def _build_left(self) -> QWidget:
        w = QWidget()
        w.setStyleSheet("background-color: #0d1117;")
        lay = QVBoxLayout(w)
        lay.setContentsMargins(16, 16, 8, 8)
        lay.setSpacing(12)

        # Drop zone
        self._drop_zone = DropZoneWidget()
        self._drop_zone.file_dropped.connect(self._on_file_dropped)
        lay.addWidget(self._drop_zone)

        # Section header
        text_hdr = QHBoxLayout()
        sec_lbl = QLabel("VĂN BẢN")
        sec_lbl.setObjectName("heading")
        self._char_lbl = QLabel("0 ký tự · 0 câu")
        self._char_lbl.setObjectName("badge")
        text_hdr.addWidget(sec_lbl)
        text_hdr.addStretch()
        text_hdr.addWidget(self._char_lbl)
        lay.addLayout(text_hdr)

        # Text editor
        self._editor = QPlainTextEdit()
        self._editor.setPlaceholderText(
            "Kéo file vào ô trên, hoặc nhập văn bản trực tiếp tại đây...\n\n"
            "Mỗi dòng / câu sẽ được xử lý riêng biệt khi xuất MP3."
        )
        self._editor.textChanged.connect(self._on_text_changed)
        lay.addWidget(self._editor, 1)

        return w

    # ── Right panel ──────────────────────────────────────────────────────

    def _build_right(self) -> QScrollArea:
        container = QWidget()
        container.setStyleSheet("background-color: #0d1117;")

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setWidget(container)
        scroll.setStyleSheet("QScrollArea { border:none; background:#0d1117; }")
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)  # Tắt scrollbar ngang

        lay = QVBoxLayout(container)
        lay.setContentsMargins(8, 16, 16, 16)
        lay.setSpacing(14)

        lay.addWidget(self._build_card_voice())
        lay.addWidget(self._build_card_whisper())
        lay.addWidget(self._build_card_export())
        lay.addWidget(self._build_card_skip_to_slide())
        lay.addStretch()

        return scroll

    # ── Card: Voice ──────────────────────────────────────────────────────

    def _build_card_voice(self) -> QFrame:
        card = QFrame()
        card.setObjectName("card")
        lay = QVBoxLayout(card)
        lay.setContentsMargins(14, 14, 14, 14)
        lay.setSpacing(8)

        # Heading
        h_row = QHBoxLayout()
        h = QLabel("🎵  GIỌNG ĐỌC")
        h.setObjectName("heading")
        self._model_status_badge = QLabel("Chưa tải")
        self._model_status_badge.setObjectName("badge-orange")
        h_row.addWidget(h)
        h_row.addStretch()
        h_row.addWidget(self._model_status_badge)
        lay.addLayout(h_row)

        # Model combo
        lay.addWidget(self._make_label("Model:"))
        self._model_combo = QComboBox()
        for name in self._vi_models:
            self._model_combo.addItem(name)
        
        # Load last selected model
        settings = QSettings("KathTTS", "KathSlideToVideoMaker")
        last_model = settings.value("last_selected_model", None)
        if last_model and last_model in self._vi_models:
            self._model_combo.setCurrentText(last_model)

        self._model_combo.currentTextChanged.connect(self._refresh_model_status)
        lay.addWidget(self._model_combo)

        # Model description
        self._model_desc = QLabel("")
        self._model_desc.setStyleSheet("color: #7d8590; font-size: 11px; background:transparent;")
        self._model_desc.setWordWrap(True)
        lay.addWidget(self._model_desc)

        # Download button + progress
        self._download_btn = QPushButton("⬇  Tải model")
        self._download_btn.setObjectName("primary")
        self._download_btn.clicked.connect(self._start_download)
        lay.addWidget(self._download_btn)

        self._download_progress = QProgressBar()
        self._download_progress.setVisible(False)
        self._download_progress.setFixedHeight(6)
        lay.addWidget(self._download_progress)

        lay.addSpacing(4)
        sep = QFrame()
        sep.setFrameShape(QFrame.Shape.HLine)
        sep.setStyleSheet("background-color: #21262d; max-height:1px; border:none;")
        lay.addWidget(sep)
        lay.addSpacing(4)

        # Speaker combo
        lay.addWidget(self._make_label("Người đọc (Speaker ID):"))
        self._speaker_combo = QComboBox()
        lay.addWidget(self._speaker_combo)

        # Speed combo
        lay.addWidget(self._make_label("Tốc độ đọc:"))
        self._speed_combo = QComboBox()
        speeds = [
            ("0.8x", 0.8),
            ("0.9x", 0.9),
            ("1.0x (Mặc định)", 1.0),
            ("1.05x", 1.05),
            ("1.1x", 1.1),
            ("1.15x", 1.15),
            ("1.2x", 1.2),
            ("1.3x", 1.3),
            ("1.5x", 1.5),
            ("1.8x", 1.8),
            ("2.0x", 2.0),
        ]
        for label, val in speeds:
            self._speed_combo.addItem(label, val)
        
        # Load last selected speed
        settings = QSettings("KathTTS", "KathSlideToVideoMaker")
        last_speed = settings.value("last_selected_speed", "1.0x (Mặc định)")
        idx = self._speed_combo.findText(last_speed)
        if idx >= 0:
            self._speed_combo.setCurrentIndex(idx)
        else:
            self._speed_combo.setCurrentIndex(2)  # Default to 1.0x

        self._speed_combo.currentTextChanged.connect(self._save_speed_setting)
        lay.addWidget(self._speed_combo)

        note = QLabel("💡 Preview để tìm giọng Nam phù hợp.")
        note.setStyleSheet("color: #484f58; font-size: 11px; background:transparent;")
        lay.addWidget(note)

        # Preview button
        self._preview_btn = QPushButton("▶  Preview giọng")
        self._preview_btn.clicked.connect(self._start_preview)
        lay.addWidget(self._preview_btn)

        return card

    # ── Card: Whisper ────────────────────────────────────────────────────

    def _build_card_whisper(self) -> QFrame:
        card = QFrame()
        card.setObjectName("card")
        lay = QVBoxLayout(card)
        lay.setContentsMargins(14, 14, 14, 14)
        lay.setSpacing(8)

        h = QLabel("🕐  WORD TIMESTAMPS")
        h.setObjectName("heading")
        lay.addWidget(h)

        self._whisper_check = QCheckBox("Phân tích timestamps từng từ (Whisper)")
        self._whisper_check.setChecked(True)
        lay.addWidget(self._whisper_check)

        desc = QLabel(
            "Dùng Whisper tiny (~75MB) để xác định thời điểm chính xác\n"
            "từng từ trong audio — cần thiết cho slide timing."
        )
        desc.setStyleSheet("color: #7d8590; font-size: 11px; background:transparent;")
        desc.setWordWrap(True)
        lay.addWidget(desc)

        return card

    # ── Card: Export ─────────────────────────────────────────────────────

    def _build_card_export(self) -> QFrame:
        card = QFrame()
        card.setObjectName("card")
        lay = QVBoxLayout(card)
        lay.setContentsMargins(14, 14, 14, 14)
        lay.setSpacing(8)

        h = QLabel("📤  XUẤT FILE")
        h.setObjectName("heading")
        lay.addWidget(h)

        # Output path row
        out_row = QHBoxLayout()
        self._out_path = QLineEdit()
        self._out_path.setPlaceholderText("Chọn đường dẫn lưu file .mp3...")
        
        # Default save path: Downloads, filename: dd-mm-yyyy.mp3
        settings = QSettings("KathTTS", "KathSlideToVideoMaker")
        last_dir = settings.value("last_export_dir", None)
        if not last_dir:
            last_dir = str(Path.home() / "Downloads")
            if not Path(last_dir).exists():
                last_dir = str(Path.home())
        
        from datetime import datetime
        default_name = f"{datetime.now().strftime('%d-%m-%Y')}.mp3"
        default_path = str(Path(last_dir) / default_name)
        self._out_path.setText(default_path)

        browse_btn = QPushButton("…")
        browse_btn.setFixedWidth(36)
        browse_btn.clicked.connect(self._browse_output)
        out_row.addWidget(self._out_path, 1)
        out_row.addWidget(browse_btn)
        lay.addLayout(out_row)

        out_note = QLabel("File .json timestamps sẽ được tạo cùng thư mục với .mp3")
        out_note.setStyleSheet("color: #484f58; font-size: 11px; background:transparent;")
        out_note.setWordWrap(True)
        lay.addWidget(out_note)

        lay.addSpacing(6)

        # Export button
        self._export_btn = QPushButton("🎵  Xuất MP3")
        self._export_btn.setObjectName("success")
        self._export_btn.clicked.connect(self._start_export)
        lay.addWidget(self._export_btn)

        return card

    # ── Card: Skip to slide (already have MP3) ────────────────────────

    def _build_card_skip_to_slide(self) -> QFrame:
        card = QFrame()
        card.setObjectName("card")
        card.setStyleSheet("""
            QFrame#card {
                background-color: #0d1f12;
                border: 1px solid #238636;
                border-radius: 10px;
                padding: 4px;
            }
        """)
        lay = QVBoxLayout(card)
        lay.setContentsMargins(14, 12, 14, 12)
        lay.setSpacing(8)

        h_row = QHBoxLayout()
        icon = QLabel("📌")
        icon.setStyleSheet("font-size:16px; background:transparent;")
        title = QLabel("ĐÃ CÓ FILE MP3?")
        title.setObjectName("heading")
        title.setStyleSheet(
            "font-size:11px; font-weight:700; color:#3fb950; "
            "letter-spacing:0.08em; background:transparent;"
        )
        h_row.addWidget(icon)
        h_row.addSpacing(6)
        h_row.addWidget(title)
        h_row.addStretch()
        lay.addLayout(h_row)

        desc = QLabel(
            "Nếu bạn đã có file .mp3 từ lần trước,\n"
            "có thể bỏ qua bước tạo MP3 và sang\n"
            "thẳng bước Đồng bộ Slide."
        )
        desc.setStyleSheet("color:#7d8590; font-size:11px; background:transparent;")
        desc.setWordWrap(True)
        lay.addWidget(desc)

        # JSON note
        json_note = QLabel(
            "💡 Lưu ý: cần file .json timestamps đi kèm MP3 để xuất video chính xác."
        )
        json_note.setStyleSheet(
            "color:#f0883e; font-size:10px; background:transparent;"
        )
        json_note.setWordWrap(True)
        lay.addWidget(json_note)

        skip_btn = QPushButton("➡  Dùng MP3 có sẵn → Đồng bộ Slide")
        skip_btn.setObjectName("success")
        skip_btn.setStyleSheet("""
            QPushButton {
                background-color: #1a7f37; color:#fff; border:none;
                border-radius:6px; font-size:12px; font-weight:700;
                padding:8px 14px; min-height:34px;
            }
            QPushButton:hover { background-color: #238636; }
            QPushButton:pressed { background-color: #116329; }
        """)
        skip_btn.clicked.connect(self._jump_to_slide_with_existing_mp3)
        lay.addWidget(skip_btn)

        return card

    # ── Footer (status bar) ──────────────────────────────────────────────

    def _build_footer(self) -> QFrame:
        footer = QFrame()
        footer.setFixedHeight(44)
        footer.setStyleSheet("""
            QFrame {
                background-color: #161b22;
                border-top: 1px solid #21262d;
            }
        """)
        lay = QHBoxLayout(footer)
        lay.setContentsMargins(16, 0, 16, 0)
        lay.setSpacing(12)

        self._status_lbl = QLabel("Sẵn sàng.")
        self._status_lbl.setStyleSheet("color: #7d8590; background:transparent;")

        self._export_progress = QProgressBar()
        self._export_progress.setFixedHeight(6)
        self._export_progress.setFixedWidth(220)
        self._export_progress.setVisible(False)

        lay.addWidget(self._status_lbl, 1)
        lay.addWidget(self._export_progress)

        return footer

    # ═══════════════════════════════════════════════════════════════════
    #  HELPERS
    # ═══════════════════════════════════════════════════════════════════

    def _make_label(self, text: str) -> QLabel:
        lbl = QLabel(text)
        lbl.setStyleSheet("color: #8b949e; font-size: 12px; background:transparent;")
        return lbl

    def _set_status(self, msg: str):
        self._status_lbl.setText(msg)

    def _refresh_model_status(self):
        model_name = self._model_combo.currentText()
        if not model_name:
            return

        # Save model selection
        settings = QSettings("KathTTS", "KathSlideToVideoMaker")
        settings.setValue("last_selected_model", model_name)

        info = self._vi_models.get(model_name, {})
        is_downloaded = self._engine.is_model_downloaded(model_name)

        # Update description
        self._model_desc.setText(info.get("description", ""))

        if is_downloaded:
            self._model_status_badge.setText("✓ Đã tải")
            self._model_status_badge.setObjectName("badge-green")
            self._download_btn.setText("✓ Đã tải")
            self._download_btn.setEnabled(False)
            self._download_btn.setObjectName("")
            self._load_speakers(model_name)
        else:
            self._model_status_badge.setText("Chưa tải")
            self._model_status_badge.setObjectName("badge-orange")
            self._download_btn.setText("⬇  Tải model")
            self._download_btn.setEnabled(True)
            self._download_btn.setObjectName("primary")
            self._speaker_combo.clear()
            self._speaker_combo.addItem("— Tải model trước —")

        # Force style refresh
        self._model_status_badge.style().unpolish(self._model_status_badge)
        self._model_status_badge.style().polish(self._model_status_badge)

    def _load_speakers(self, model_name: str):
        self._speaker_combo.clear()
        speakers = self._engine.get_speakers(model_name)
        for name, sid in speakers:
            self._speaker_combo.addItem(f"{name}  (ID: {sid})", sid)

    def _save_speed_setting(self, speed_text: str):
        settings = QSettings("KathTTS", "KathSlideToVideoMaker")
        settings.setValue("last_selected_speed", speed_text)

    # ═══════════════════════════════════════════════════════════════════
    #  EVENT HANDLERS
    # ═══════════════════════════════════════════════════════════════════

    def _on_file_dropped(self, path: str):
        try:
            text = self._read_file(path)
            self._editor.setPlainText(text)
            fname = Path(path).name
            self._set_status(f"✓ Đã tải: {fname}")
        except Exception as exc:
            QMessageBox.critical(self, "Lỗi đọc file", str(exc))

    def _on_text_changed(self):
        text = self._editor.toPlainText()
        chars = len(text)

        from app.core.mp3_exporter import split_into_sentences
        sentences = split_into_sentences(text) if text.strip() else []

        self._char_lbl.setText(f"{chars:,} ký tự  ·  {len(sentences)} câu")

    # ── Download ─────────────────────────────────────────────────────────

    def _start_download(self):
        model_name = self._model_combo.currentText()
        self._download_btn.setEnabled(False)
        self._download_progress.setVisible(True)
        self._download_progress.setValue(0)
        self._set_status("Đang tải model…")

        self._download_thread = DownloadThread(self._engine, model_name)
        self._download_thread.progress.connect(self._on_download_progress)
        self._download_thread.finished.connect(self._on_download_finished)
        self._download_thread.start()

    def _on_download_progress(self, pct: int, msg: str):
        self._download_progress.setValue(pct)
        self._set_status(f"{msg}  ({pct}%)")

    def _on_download_finished(self, success: bool, error: str):
        self._download_progress.setVisible(False)
        if success:
            self._set_status("✓ Tải model xong!")
            self._refresh_model_status()
        else:
            self._download_btn.setEnabled(True)
            self._set_status("❌ Tải model thất bại.")
            QMessageBox.critical(self, "Lỗi tải model", error)

    # ── Preview ──────────────────────────────────────────────────────────

    def _start_preview(self):
        model_name = self._model_combo.currentText()
        if not self._engine.is_model_downloaded(model_name):
            QMessageBox.warning(self, "Chưa tải model",
                                "Vui lòng tải model trước khi preview.")
            return

        speaker_id = self._speaker_combo.currentData() or 0
        speed = self._speed_combo.currentData() or 1.0
        self._preview_btn.setEnabled(False)
        self._set_status("Đang tạo preview giọng…")

        # Lấy văn bản được chọn hoặc dòng đầu tiên để preview
        cursor = self._editor.textCursor()
        preview_text = cursor.selectedText().strip()
        preview_text = preview_text.replace("\u2029", " ") # Dọn dẹp dòng xuống dòng của PyQt

        if not preview_text:
            # Nếu không bôi đen, lấy dòng đầu tiên của văn bản
            all_text = self._editor.toPlainText().strip()
            if all_text:
                lines = [line.strip() for line in all_text.split("\n") if line.strip()]
                if lines:
                    preview_text = lines[0]

        if not preview_text:
            # Fallback mặc định
            preview_text = (
                "Xin chào! Đây là giọng đọc thử nghiệm. "
                "Dự án Slide To Video sẽ dùng giọng này để tạo bài thuyết trình."
            )

        # Giới hạn độ dài preview để tránh chờ quá lâu
        if len(preview_text) > 150:
            preview_text = preview_text[:150] + "..."

        self._preview_thread = PreviewThread(
            self._engine, model_name, speaker_id, preview_text, speed=speed
        )
        self._preview_thread.finished.connect(self._on_preview_finished)
        self._preview_thread.start()

    def _on_preview_finished(self, success: bool, wav_path: str, error: str):
        self._preview_btn.setEnabled(True)
        if success:
            self._set_status("▶ Đang phát preview…")
            try:
                winsound.PlaySound(wav_path, winsound.SND_FILENAME | winsound.SND_ASYNC)
            except Exception:
                # Fallback: dùng os.startfile
                import os
                os.startfile(wav_path)
        else:
            self._set_status("❌ Preview thất bại.")
            QMessageBox.critical(self, "Lỗi preview", error)

    # ── Export ───────────────────────────────────────────────────────────

    def _browse_output(self):
        settings = QSettings("KathTTS", "KathSlideToVideoMaker")
        current_text = self._out_path.text().strip()
        
        initial_dir = ""
        if current_text:
            try:
                initial_dir = str(Path(current_text).parent)
            except Exception:
                pass
                
        if not initial_dir or not Path(initial_dir).exists():
            initial_dir = settings.value("last_export_dir", None)
            
        if not initial_dir or not Path(initial_dir).exists():
            initial_dir = str(Path.home() / "Downloads")
            if not Path(initial_dir).exists():
                initial_dir = str(Path.home())

        from datetime import datetime
        default_name = f"{datetime.now().strftime('%d-%m-%Y')}.mp3"
        default_path = str(Path(initial_dir) / default_name)

        path, _ = QFileDialog.getSaveFileName(
            self, "Chọn vị trí lưu file MP3", default_path,
            "MP3 Audio (*.mp3)"
        )
        if path:
            if not path.lower().endswith(".mp3"):
                path += ".mp3"
            self._out_path.setText(path)
            selected_dir = str(Path(path).parent)
            settings.setValue("last_export_dir", selected_dir)

    def _start_export(self):
        text = self._editor.toPlainText().strip()
        if not text:
            QMessageBox.warning(self, "Thiếu văn bản",
                                "Vui lòng nhập hoặc tải văn bản trước khi xuất.")
            return

        output_path = self._out_path.text().strip()
        if not output_path:
            QMessageBox.warning(self, "Thiếu đường dẫn",
                                "Vui lòng chọn đường dẫn lưu file MP3.")
            return

        model_name = self._model_combo.currentText()
        if not self._engine.is_model_downloaded(model_name):
            QMessageBox.warning(self, "Chưa tải model",
                                "Vui lòng tải model TTS trước khi xuất.")
            return

        # Load model
        try:
            self._engine.load_model(model_name)
        except Exception as exc:
            QMessageBox.critical(self, "Lỗi load model", str(exc))
            return

        speaker_id  = self._speaker_combo.currentData() or 0
        use_whisper = self._whisper_check.isChecked()
        speed       = self._speed_combo.currentData() or 1.0

        self._export_btn.setEnabled(False)
        self._export_progress.setVisible(True)
        self._export_progress.setValue(0)

        self._export_thread = ExportThread(
            self._pipeline, text, self._engine,
            speaker_id, output_path, use_whisper, speed=speed,
        )
        self._export_thread.progress.connect(self._on_export_progress)
        self._export_thread.finished.connect(self._on_export_finished)
        self._export_thread.start()

    def _on_export_progress(self, pct: int, msg: str):
        self._export_progress.setValue(pct)
        self._set_status(msg)

    def _on_export_finished(self, success: bool, error: str):
        self._export_progress.setVisible(False)
        self._export_btn.setEnabled(True)

        if success:
            mp3_path = self._out_path.text()
            self._set_status("✓ Xuất thành công! Đang chuyển sang Đồng bộ Slide…")
            # Auto-switch: JSON luôn có vì vừa xuất xong
            self._switch_to_slide_tab(mp3_path, has_json=True)
        else:
            self._set_status("❌ Xuất thất bại.")
            QMessageBox.critical(self, "Lỗi xuất file", error)

    # ── Tab switching ─────────────────────────────────────────────────────

    def _jump_to_slide_with_existing_mp3(self):
        """
        Cho phép nhảy sang Tab 2 khi người dùng đã có file MP3 từ trước.
        Bước 1: Chọn file MP3.
        Bước 2: Tự tìm JSON cùng tên. Nếu không thấy → hỏi browse JSON riêng.
        """
        # ── Bước 1: Chọn file MP3 ──────────────────────────────────────
        settings = QSettings("KathTTS", "KathSlideToVideoMaker")
        last_dir = settings.value("last_export_dir", None)
        if not last_dir or not Path(last_dir).exists():
            last_dir = str(Path.home() / "Downloads")
            if not Path(last_dir).exists():
                last_dir = str(Path.home())

        mp3_path, _ = QFileDialog.getOpenFileName(
            self,
            "Bước 1/2 — Chọn file MP3",
            last_dir,
            "MP3 Audio (*.mp3);;Tất cả (*.*)",
        )
        if not mp3_path:
            return

        # ── Bước 2: Tự động tìm JSON cùng tên cùng thư mục ───────────
        auto_json = Path(mp3_path).with_suffix(".json")
        has_json  = auto_json.exists()

        if not has_json:
            # Hỏi user muốn chỉ định JSON ở nơi khác, hay bỏ qua
            from PyQt6.QtWidgets import QMessageBox as MB
            reply = MB.question(
                self,
                "Bước 2/2 — Tìm file timestamps",
                f"Không tìm thấy file .json cùng tên:\n"
                f"  {auto_json}\n\n"
                "Bạn muốn làm gì?",
                MB.StandardButton.Open   |   # "Chọn file JSON…"
                MB.StandardButton.Ignore |   # "Bỏ qua, tiếp tục"
                MB.StandardButton.Cancel,    # "Huỷ"
                MB.StandardButton.Open,
            )

            if reply == MB.StandardButton.Cancel:
                return

            elif reply == MB.StandardButton.Open:
                # Cho browse JSON thủ công
                json_path, _ = QFileDialog.getOpenFileName(
                    self,
                    "Bước 2/2 — Chọn file timestamps JSON",
                    str(Path(mp3_path).parent),  # mở cùng thư mục MP3
                    "JSON Timestamps (*.json);;Tất cả (*.*)",
                )
                if not json_path:
                    return  # user cancel dialog JSON
                # Copy/link json vào cùng tên với mp3 (không copy, chỉ track path)
                # → truyền has_json=True vì user đã chỉ định file
                # Lưu đường dẫn json custom để dùng sau
                self._custom_json_path = json_path
                has_json = True
            else:
                # Ignore — tiếp tục không có JSON
                self._custom_json_path = ""
                has_json = False
        else:
            self._custom_json_path = ""  # dùng auto-detect

        # ── Bước 3: Tìm kịch bản văn bản ─────────────────────────────
        script_text = self._editor.toPlainText().strip()

        # Ưu tiên đọc từ file .slides.json (đã lưu dự án trước đó)
        slides_json = Path(mp3_path).with_suffix(".slides.json")
        resolved_json = self._custom_json_path if (not has_json or not auto_json.exists()) else str(auto_json)
        if not script_text and slides_json.exists():
            try:
                import json as _json
                with open(slides_json, "r", encoding="utf-8") as f:
                    data = _json.load(f)
                script_text = data.get("script_text", "")
            except Exception:
                pass

        # Tiếp theo: reconstruct từ timestamps JSON (sentences[].text)
        if not script_text and resolved_json and Path(resolved_json).exists():
            try:
                import json as _json
                with open(resolved_json, "r", encoding="utf-8") as f:
                    json_data = _json.load(f)
                sentences = [s.get("text", "") for s in json_data.get("sentences", [])]
                script_text = " ".join(sentences)
            except Exception:
                pass

        # Fallback: hỏi user chọn file kịch bản
        if not script_text:
            reply = QMessageBox.question(
                self,
                "Kịch bản trống",
                "Không tìm thấy kịch bản văn bản tự động.\n\n"
                "Bạn có muốn chọn file kịch bản (.txt / .docx) không?\n"
                "(Hoặc chọn Không để tiếp tục với editor trống)",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            )
            if reply == QMessageBox.StandardButton.Yes:
                txt_path, _ = QFileDialog.getOpenFileName(
                    self, "Chọn file kịch bản", str(Path(mp3_path).parent),
                    "Văn bản (*.txt *.docx);;Tất cả (*.*)"
                )
                if txt_path:
                    try:
                        script_text = _read_script_file(txt_path)
                    except Exception as e:
                        QMessageBox.warning(self, "Lỗi đọc file", str(e))

        self._switch_to_slide_tab(mp3_path, script_text=script_text, has_json=has_json)



    def _switch_to_slide_tab(self, mp3_path: str, script_text: str = "", has_json: bool = True):
        """Chuyển sang tab đồng bộ slide và truyền context."""
        if not script_text:
            script_text = self._editor.toPlainText()
        self._slide_sync_tab.load_context(script_text, mp3_path, has_json=has_json)
        self._stack.setCurrentIndex(1)
        # Update step indicator
        self._step1_lbl.setText("✓ Bước 1: Tạo MP3")
        self._step1_lbl.setObjectName("step-inactive")
        self._step2_lbl.setText("● Bước 2: Đồng bộ Slide")
        self._step2_lbl.setObjectName("step-active")
        # Force style refresh
        for lbl in [self._step1_lbl, self._step2_lbl]:
            lbl.style().unpolish(lbl)
            lbl.style().polish(lbl)

    def _go_to_mp3_tab(self):
        """Quay lại tab tạo MP3."""
        self._stack.setCurrentIndex(0)
        self._step1_lbl.setText("● Bước 1: Tạo MP3")
        self._step1_lbl.setObjectName("step-active")
        self._step2_lbl.setText("○ Bước 2: Đồng bộ Slide")
        self._step2_lbl.setObjectName("step-inactive")
        for lbl in [self._step1_lbl, self._step2_lbl]:
            lbl.style().unpolish(lbl)
            lbl.style().polish(lbl)

    def _on_export_video_requested(self, slides: list):
        """Xử lý khi người dùng nhấn 'Tiếp tục → Xuất Video'."""
        # Cập nhật step 3 indicator
        self._step2_lbl.setText("✓ Bước 2: Đồng bộ Slide")
        self._step2_lbl.setObjectName("step-inactive")
        self._step3_lbl.setText("● Bước 3: Xuất Video")
        self._step3_lbl.setObjectName("step-active")
        for lbl in [self._step2_lbl, self._step3_lbl]:
            lbl.style().unpolish(lbl)
            lbl.style().polish(lbl)
        QMessageBox.information(
            self, "Bước 3: Xuất Video",
            f"Sẵn sàng xuất video với {len(slides)} slide.\n\n"
            "Tính năng Xuất Video sẽ được thêm vào ở phiên bản tiếp theo."
        )
